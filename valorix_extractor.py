#!/usr/bin/env python3
"""
VALORIX360 · Extractor de estados financieros (PDF) -> plantilla de diagnóstico

Uso:
    python valorix_extractor.py estados.pdf \
        --plantilla valorix360_plantilla_diagnostico.xlsx \
        --salida diagnostico_cliente.xlsx

Flujo:
  1. Lee el PDF (texto nativo; si una página no tiene texto aplica OCR con Tesseract).
  2. Detecta secciones: balance, resultados, flujo de efectivo y notas.
  3. Convierte cifras (formato colombiano, paréntesis = negativo) y las mapea a un catálogo estándar.
  4. Valida cuadres contables y deja marcado todo lo dudoso.
  5. Llena la hoja Datos de la plantilla y agrega una hoja "Revisión" y una hoja "Notas".

IMPORTANTE: es semiautomático. Revisa la hoja "Revisión" antes de usar los resultados.
Los estados financieros son información confidencial del cliente: procésalos solo localmente
o con herramientas que cumplan el acuerdo de confidencialidad firmado.
"""
import argparse
import difflib
import json
import re
import sys
import unicodedata
from datetime import date

import pdfplumber
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

# --------------------------------------------------------------------------------------
# 1. CATÁLOGO ESTÁNDAR (editable: aquí se amplía el diccionario con cada cliente nuevo)
# --------------------------------------------------------------------------------------
# sec   : secciones donde se busca la cuenta (bal=balance, er=resultados, cf=flujo de efectivo)
# grupo : grupo del balance donde debe estar (ac, anc, pc, pnc, pat). None = cualquiera
# exact : solo coincidencia exacta (totales y cuentas genéricas)
# suma  : si hay varias líneas con la misma clave, se suman
CATALOGO = {
    # ---- Balance
    "efectivo": dict(sec=["bal"], grupo=["ac"], syn=[
        "efectivo y equivalentes de efectivo", "efectivo y equivalentes", "efectivo", "caja y bancos", "disponible"]),
    "cxc": dict(sec=["bal"], grupo=["ac"], syn=[
        "deudores comerciales", "deudores comerciales y otras cuentas por cobrar", "cuentas por cobrar comerciales",
        "cuentas comerciales por cobrar", "cuentas por cobrar", "clientes", "deudores"]),
    "inventarios": dict(sec=["bal"], grupo=["ac"], syn=["inventarios", "inventario", "existencias"]),
    "capital": dict(sec=["bal"], grupo=["pat"], exact=True, syn=[
        "capital social", "capital suscrito y pagado", "capital emitido", "capital social y primas", "capital"]),
    "total_ac": dict(sec=["bal"], exact=True, syn=[
        "total activo corriente", "total activos corrientes", "total activo circulante"]),
    "total_anc": dict(sec=["bal"], exact=True, syn=[
        "total activo no corriente", "total activos no corrientes"]),
    "total_activos": dict(sec=["bal"], exact=True, syn=["total activos", "total activo"]),
    "oblig_cp": dict(sec=["bal"], grupo=["pc"], syn=[
        "obligaciones financieras", "obligaciones financieras corrientes", "prestamos y obligaciones",
        "pasivos financieros", "prestamos", "creditos bancarios", "prestamos y obligaciones financieras"]),
    "oblig_lp": dict(sec=["bal"], grupo=["pnc"], syn=[
        "obligaciones financieras", "obligaciones financieras no corrientes", "prestamos y obligaciones",
        "pasivos financieros", "prestamos", "creditos bancarios", "prestamos y obligaciones financieras"]),
    "proveedores": dict(sec=["bal"], grupo=["pc"], syn=[
        "proveedores", "cuentas comerciales por pagar", "cuentas por pagar comerciales", "acreedores comerciales",
        "cuentas por pagar a proveedores", "cuentas por pagar comerciales y otras cuentas por pagar",
        "cuentas por pagar"]),
    "total_pc": dict(sec=["bal"], exact=True, syn=["total pasivo corriente", "total pasivos corrientes"]),
    "total_pnc": dict(sec=["bal"], exact=True, syn=["total pasivo no corriente", "total pasivos no corrientes"]),
    "total_pasivos": dict(sec=["bal"], exact=True, syn=["total pasivos", "total pasivo"]),
    "patrimonio": dict(sec=["bal"], exact=True, syn=[
        "total patrimonio", "patrimonio total", "total patrimonio neto", "total patrimonio de los accionistas"]),
    "total_pas_pat": dict(sec=["bal"], exact=True, syn=[
        "total pasivos y patrimonio", "total pasivo y patrimonio"]),
    # ---- Estado de resultados
    "ingresos": dict(sec=["er"], syn=[
        "ingresos de actividades ordinarias", "ingresos operacionales", "ventas netas", "ingresos por ventas",
        "ingresos ordinarios", "ventas", "ingresos"]),
    "costo": dict(sec=["er"], syn=[
        "costo de ventas", "costos de ventas", "costo de ventas y prestacion de servicios",
        "costo de mercancia vendida", "costo de los bienes vendidos"]),
    "utilidad_bruta": dict(sec=["er"], exact=True, syn=["utilidad bruta", "ganancia bruta", "resultado bruto"]),
    "gastos_op": dict(sec=["er"], suma=True, syn=[
        "gastos de administracion", "gastos de ventas", "gastos de distribucion",
        "gastos de administracion y ventas", "gastos operacionales", "otros gastos operacionales"]),
    "utilidad_operativa": dict(sec=["er"], exact=True, syn=[
        "utilidad operativa", "ganancia operativa", "resultado operacional", "utilidad operacional"]),
    "gastos_fin": dict(sec=["er"], suma=True, syn=[
        "gastos financieros", "costos financieros", "gastos por intereses"]),
    "impuestos": dict(sec=["er"], suma=True, syn=[
        "gasto por impuesto a las ganancias", "impuesto de renta", "impuesto sobre la renta", "gasto por impuesto",
        "impuesto a las ganancias", "provision impuesto de renta", "gasto por impuesto de renta"]),
    "utilidad_neta": dict(sec=["er"], exact=True, syn=[
        "utilidad neta", "ganancia del periodo", "resultado del periodo", "utilidad del ejercicio",
        "ganancia neta", "resultado neto", "utilidad del periodo"]),
    # ---- Flujo de efectivo
    "cfo": dict(sec=["cf"], syn=[
        "flujo de efectivo neto de actividades de operacion",
        "flujos de efectivo netos procedentes de actividades de operacion",
        "efectivo neto provisto por las actividades de operacion",
        "efectivo generado por actividades de operacion",
        "flujo de efectivo generado por actividades de operacion",
        "efectivo neto generado por actividades de operacion"]),
    "da": dict(sec=["cf"], suma=True, syn=[
        "depreciacion y amortizacion", "depreciacion", "amortizacion", "gasto por depreciacion y amortizacion",
        "depreciacion de propiedades planta y equipo"]),
    "capex": dict(sec=["cf"], suma=True, syn=[
        "adquisicion de propiedades planta y equipo", "compra de propiedades planta y equipo",
        "compras de propiedad planta y equipo", "adquisiciones de propiedad planta y equipo",
        "inversion en propiedad planta y equipo", "adquisicion de activos fijos"]),
    "pago_capital": dict(sec=["cf"], suma=True, syn=[
        "pago de obligaciones financieras", "pagos de obligaciones financieras", "pagos de prestamos",
        "pago de capital", "reembolso de prestamos"]),
    "intereses_pagados": dict(sec=["cf"], suma=True, syn=["intereses pagados", "pago de intereses"]),
}

# Fila de cada dato en la hoja Datos de la plantilla
FILAS_DATOS = {
    "efectivo": 7, "cxc": 8, "inventarios": 9, "otros_ac": 10, "anc": 12,
    "oblig_cp": 14, "proveedores": 15, "otros_pc": 16, "pnc": 18, "patrimonio": 20,
    "ingresos": 24, "costo": 25, "gastos_op": 26, "da": 27, "gastos_fin": 28, "impuestos": 29,
    "cfo": 36, "capex": 37, "servicio_deuda": 39, "dias": 40, "deuda_lp": 43, "capital": 62,
}

# Filas de datos adicionales que el PDF no trae: se dejan para digitar a mano (se limpian los ejemplos)
FILAS_MANUALES_CERO = {44: "Gastos no recurrentes", 45: "Ingresos no recurrentes"}
FILAS_MANUALES_VACIAS = {
    47: "Ingresos del año completo anterior", 48: "Costo de ventas del año completo anterior",
    49: "Gastos operativos (sin D&A) del año completo anterior", 50: "D&A del año completo anterior",
    51: "Flujo de caja operativo del año completo anterior", 52: "Utilidad neta del año completo anterior",
    54: "Endeudamiento máximo pactado (covenant)", 55: "Cobertura mínima del servicio de la deuda",
    56: "Deuda neta / EBITDA máximo", 57: "Vencimiento de la deuda: año 1", 58: "Vencimiento de la deuda: año 2",
    59: "Vencimiento de la deuda: año 3", 60: "Vencimiento de la deuda: año 4 o más",
    63: "Costo del patrimonio (Ke)",
}

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
         "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

CATEGORIAS_NOTAS = {
    "Deuda y garantías": ["obligaciones financieras", "prestamo", "credito", "tasa", "garantia", "covenant",
                          "vencimiento", "pagare", "prenda", "hipoteca"],
    "Cartera": ["deudores", "cuentas por cobrar", "cartera", "deterioro", "edades", "incobrable"],
    "Inventarios": ["inventario", "obsolesc", "baja rotacion"],
    "Contingencias": ["contingencia", "litigio", "demanda", "proceso judicial", "proceso laboral"],
    "Partes relacionadas": ["partes relacionadas", "vinculados", "accionistas", "junta directiva"],
    "Hechos posteriores": ["hechos posteriores", "eventos posteriores"],
    "Continuidad del negocio": ["negocio en marcha", "empresa en funcionamiento", "continuidad", "incertidumbre"],
    "Impuestos": ["impuesto diferido", "impuesto de renta", "tasa efectiva"],
}

# --------------------------------------------------------------------------------------
# 2. UTILIDADES DE TEXTO Y NÚMEROS
# --------------------------------------------------------------------------------------
def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


NUM_RE = re.compile(r"^[\(\-]*\$?\s*[\(\-]*\d[\d.,]*\)?%?$")


def es_numero(tok: str) -> bool:
    return bool(NUM_RE.match(tok.strip()))


def parse_numero(tok: str):
    """Convierte '(1.234.567)' -> -1234567.0 ; '1.234,50' -> 1234.5 ; '1,234.50' -> 1234.5"""
    t = tok.strip().replace("$", "").replace(" ", "")
    neg = t.startswith("(") or t.startswith("-") or t.endswith(")")
    t = t.strip("()-%")
    if not t:
        return None
    if "." in t and "," in t:
        dec = "," if t.rfind(",") > t.rfind(".") else "."
        mil = "." if dec == "," else ","
        t = t.replace(mil, "").replace(dec, ".")
    elif "." in t:
        t = t.replace(".", "") if re.fullmatch(r"\d{1,3}(\.\d{3})+", t) else t
    elif "," in t:
        t = t.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", t) else t.replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def separar_linea(linea: str, n_cols: int):
    """Devuelve (etiqueta, [valores]) o None. Ignora la columna 'Nota' (entero corto)."""
    toks = linea.split()
    cola = []
    while toks and es_numero(toks[-1]):
        cola.insert(0, toks.pop())
    if len(cola) < n_cols or not toks:
        return None
    valores = cola[-n_cols:]
    extra = cola[:-n_cols]
    # cualquier número que sobre y sea un entero corto es la referencia a la nota
    if extra and not (len(extra) == 1 and re.fullmatch(r"\d{1,2}", extra[0])):
        toks += extra  # parte de la etiqueta (p. ej. "Año 2025")
    vals = [parse_numero(v) for v in valores]
    if any(v is None for v in vals):
        return None
    return " ".join(toks), vals


def anios_en_linea(linea: str):
    ys = re.findall(r"\b(20\d{2})\b", linea)
    ys = [int(y) for y in ys]
    return ys if len(set(ys)) >= 2 else None


# --------------------------------------------------------------------------------------
# 3. LECTURA DEL PDF (texto nativo + OCR de respaldo)
# --------------------------------------------------------------------------------------
def leer_paginas(ruta_pdf, avisos):
    paginas = []
    with pdfplumber.open(ruta_pdf) as pdf:
        for i, pg in enumerate(pdf.pages, 1):
            texto = pg.extract_text() or ""
            ocr = False
            if len(texto.strip()) < 40:
                try:
                    import pytesseract
                    img = pg.to_image(resolution=250).original
                    try:
                        texto = pytesseract.image_to_string(img, lang="spa+eng", config="--psm 6")
                    except Exception:
                        texto = pytesseract.image_to_string(img, config="--psm 6")
                    ocr = True
                    avisos.append(f"Página {i}: sin texto nativo, se usó OCR. Verifica TODAS las cifras de esa página.")
                except Exception as e:  # noqa: BLE001
                    avisos.append(f"Página {i}: sin texto y OCR no disponible ({e}).")
            paginas.append({"n": i, "texto": texto, "ocr": ocr})
    return paginas


# --------------------------------------------------------------------------------------
# 4. MAPEO A CATÁLOGO
# --------------------------------------------------------------------------------------
def mapear(etiqueta_norm, seccion, grupo):
    """Devuelve (clave, confianza, método) o (None, 0, '')."""
    def aplica(spec):
        if seccion not in spec["sec"]:
            return False
        g = spec.get("grupo")
        return not g or grupo in g

    for clave, spec in CATALOGO.items():
        if aplica(spec) and etiqueta_norm in spec["syn"]:
            return clave, 1.0, "exacto"
    cands = []
    for clave, spec in CATALOGO.items():
        if not aplica(spec) or spec.get("exact"):
            continue
        for s in spec["syn"]:
            if len(s) >= 10 and etiqueta_norm.startswith(s):
                cands.append((len(s), clave))
    if cands:
        return sorted(cands, reverse=True)[0][1], 0.9, "prefijo"
    mejor = (0.0, None)
    for clave, spec in CATALOGO.items():
        if not aplica(spec) or spec.get("exact"):
            continue
        for s in spec["syn"]:
            r = difflib.SequenceMatcher(None, etiqueta_norm, s).ratio()
            if r > mejor[0]:
                mejor = (r, clave)
    if mejor[0] >= 0.9:
        return mejor[1], round(mejor[0], 2), "aproximado"
    return None, 0.0, ""


def detectar_grupo(txt_norm):
    if "no corriente" in txt_norm or "no corrientes" in txt_norm:
        if txt_norm.startswith("activo"):
            return "anc"
        if txt_norm.startswith("pasivo"):
            return "pnc"
    if txt_norm.startswith("activo") and "corriente" in txt_norm:
        return "ac"
    if txt_norm.startswith("pasivo") and "corriente" in txt_norm:
        return "pc"
    if txt_norm.startswith("patrimonio"):
        return "pat"
    return None


def detectar_seccion(txt_norm):
    if "estado de situacion financiera" in txt_norm or "balance general" in txt_norm:
        return "bal"
    if "estado de resultados" in txt_norm or "estado del resultado" in txt_norm:
        return "er"
    if "flujos de efectivo" in txt_norm or "flujo de efectivo" in txt_norm:
        return "cf"
    if "cambios en el patrimonio" in txt_norm:
        return "otro"
    if txt_norm in ("notas a los estados financieros", "notas explicativas") or \
            txt_norm.startswith("notas a los estados financieros"):
        return "notas"
    return None


# --------------------------------------------------------------------------------------
# 5. EXTRACCIÓN PRINCIPAL
# --------------------------------------------------------------------------------------
def extraer(ruta_pdf):
    avisos = []
    paginas = leer_paginas(ruta_pdf, avisos)
    texto_total = norm(" ".join(p["texto"][:1500] for p in paginas[:6]))

    log, sin_mapear = [], []
    datos = {}            # clave -> {"vals":[v0,v1], "years":[y0,y1], "lineas":[..]}
    notas_lineas = []
    seccion, grupo, years = None, None, None
    fecha_er = None

    for p in paginas:
        for linea in p["texto"].splitlines():
            linea = linea.strip()
            if not linea:
                continue
            ln = norm(linea)
            tiene_num = any(es_numero(t) for t in linea.split()[-2:])

            # secciones (solo en líneas sin cifras)
            if not tiene_num or anios_en_linea(linea):
                sec = detectar_seccion(ln)
                if sec:
                    if sec != seccion:
                        years = None
                    seccion, grupo = sec, None
                    continue
            if seccion == "notas":
                notas_lineas.append(linea)
                continue
            if seccion in (None, "otro"):
                continue

            # encabezado de columnas con años
            ys = anios_en_linea(linea)
            if ys and not tiene_num or (ys and len(linea.split()) > 4 and not re.search(r"\d{1,3}(\.\d{3}){1,}", linea)):
                years = ys[:2]
                if seccion == "er" and fecha_er is None:
                    m = re.search(r"(\d{1,2}) de (" + "|".join(MESES) + r") de (20\d{2})", linea.lower())
                    if m:
                        fecha_er = date(int(m.group(3)), MESES[m.group(2)], int(m.group(1)))
                continue

            # encabezados de grupo del balance
            if seccion == "bal" and not tiene_num:
                g = detectar_grupo(ln)
                if g:
                    grupo = g
                continue

            sep = separar_linea(linea, 2)
            if not sep:
                continue
            etiqueta, vals = sep
            en = norm(etiqueta)
            if not en or en in ("nota", "notas"):
                continue
            clave, conf, metodo = mapear(en, seccion, grupo)
            registro = dict(pagina=p["n"], seccion=seccion, grupo=grupo or "", etiqueta=etiqueta,
                            valores=vals, clave=clave or "", confianza=conf, metodo=metodo, ocr=p["ocr"])
            if not clave:
                sin_mapear.append(registro)
                continue
            spec = CATALOGO[clave]
            y = years or [None, None]
            if clave in datos and not spec.get("suma"):
                registro["metodo"] = "duplicado (ignorado)"
                log.append(registro)
                continue
            if clave in datos:
                datos[clave]["vals"] = [a + b for a, b in zip(datos[clave]["vals"], vals)]
                datos[clave]["lineas"].append(etiqueta)
            else:
                datos[clave] = dict(vals=list(vals), years=y, lineas=[etiqueta], conf=conf, ocr=p["ocr"])
            datos[clave]["conf"] = min(datos[clave]["conf"], conf)
            datos[clave]["ocr"] = datos[clave]["ocr"] or p["ocr"]
            log.append(registro)

    # unidades
    if "millones de pesos" in texto_total or "cifras en millones" in texto_total:
        factor, unidad = 1.0, "millones de pesos"
    elif "miles de pesos" in texto_total or "cifras en miles" in texto_total or "expresado en miles" in texto_total:
        factor, unidad = 1e-3, "miles de pesos"
    elif "pesos colombianos" in texto_total or "expresado en pesos" in texto_total or "en pesos" in texto_total:
        factor, unidad = 1e-6, "pesos"
    else:
        factor, unidad = None, None
        avisos.append("No se detectó la unidad (pesos / miles / millones). Usa --unidad para indicarla.")

    # columna actual / anterior según los años
    ref_years = next((d["years"] for d in datos.values() if d["years"][0] and d["years"][1]), [None, None])
    if ref_years[0] and ref_years[1]:
        i_act = 0 if ref_years[0] >= ref_years[1] else 1
    else:
        i_act = 0
        avisos.append("No se detectaron los años de las columnas; se asumió que la 1.ª columna es el periodo actual.")
    return dict(datos=datos, log=log, sin_mapear=sin_mapear, avisos=avisos, factor=factor, unidad=unidad,
                i_act=i_act, years=ref_years, fecha_er=fecha_er, notas_lineas=notas_lineas)


# --------------------------------------------------------------------------------------
# 6. CONSOLIDAR, DERIVAR Y VALIDAR
# --------------------------------------------------------------------------------------
def consolidar(res, factor_override=None, dias_act=None, dias_ant=None, da_en_gastos=True):
    d = res["datos"]
    factor = factor_override if factor_override is not None else res["factor"]
    if factor is None:
        factor = 1.0
    ia, ip = res["i_act"], 1 - res["i_act"]

    def get(clave, idx):
        if clave not in d:
            return None
        return d[clave]["vals"][idx] * factor

    out = {}      # clave -> [actual, anterior]
    meta = {}     # clave -> nota/origen
    for k in ("efectivo", "cxc", "inventarios", "oblig_cp", "proveedores", "patrimonio",
              "ingresos", "costo", "gastos_fin", "impuestos", "cfo", "da"):
        out[k] = [get(k, ia), get(k, ip)]
        if k in d and d[k]["conf"] < 1.0:
            meta[k] = f"Mapeo {d[k]['conf']:.0%} de confianza: {', '.join(d[k]['lineas'])}"
        if k in d and d[k]["ocr"]:
            meta[k] = "OCR: verificar contra el PDF"
    # impuestos/gastos financieros pueden venir negativos (como gasto): trabajar en positivo
    for k in ("gastos_fin", "impuestos", "costo"):
        out[k] = [abs(v) if v is not None else None for v in out[k]]

    def derivar(nombre, fn, nota):
        vals = []
        for idx in (ia, ip):
            try:
                v = fn(idx)
            except TypeError:
                v = None
            vals.append(v)
        out[nombre] = vals
        meta[nombre] = nota

    derivar("otros_ac", lambda i: get("total_ac", i) - (get("efectivo", i) or 0) - (get("cxc", i) or 0)
            - (get("inventarios", i) or 0), "Derivado: total activo corriente − efectivo − cartera − inventarios")
    derivar("anc", lambda i: get("total_anc", i), "Total activo no corriente del PDF")
    derivar("otros_pc", lambda i: get("total_pc", i) - (get("oblig_cp", i) or 0) - (get("proveedores", i) or 0),
            "Derivado: total pasivo corriente − obligaciones financieras − proveedores")
    derivar("pnc", lambda i: get("total_pnc", i) if "total_pnc" in d else get("total_pasivos", i) - get("total_pc", i),
            "Total pasivo no corriente del PDF")

    # gastos operativos sin depreciación
    gop = []
    for idx in (ia, ip):
        g = get("gastos_op", idx)
        if g is None:
            gop.append(None)
            continue
        g = abs(g)
        da = abs(get("da", idx) or 0)
        gop.append(g - da if da_en_gastos else g)
    out["gastos_op"] = gop
    meta["gastos_op"] = ("Gastos de administración + ventas" + (" − depreciación y amortización (del flujo de efectivo): "
                         "se asume que la D&A está dentro de los gastos. Verifica en las notas."
                         if da_en_gastos else " (sin restar D&A)"))
    out["da"] = [abs(v) if v is not None else None for v in out["da"]]

    out["capex"] = [abs(get("capex", i)) if get("capex", i) is not None else None for i in (ia, ip)]

    out["capital"] = [get("capital", ia), get("capital", ip)]
    meta["capital"] = ("Capital social del PDF" if "capital" in d else
                       "No se encontró el capital social: digítalo a mano (se usa para estimar utilidades retenidas)")

    lp = [get("oblig_lp", i) for i in (ia, ip)]
    out["deuda_lp"] = [v if v is not None else 0.0 for v in lp]
    meta["deuda_lp"] = ("Obligaciones financieras no corrientes del PDF" if any(v is not None for v in lp)
                        else "No se encontró deuda no corriente: se asumió 0. Verifica en las notas.")

    sd = []
    for idx in (ia, ip):
        cap, inte = get("pago_capital", idx), get("intereses_pagados", idx)
        sd.append(None if cap is None and inte is None else abs(cap or 0) + abs(inte or 0))
    out["servicio_deuda"] = sd
    meta["servicio_deuda"] = ("Pagos de capital + intereses pagados (flujo de efectivo)"
                              if any(v is not None for v in sd) else "No encontrado en el flujo de efectivo")

    # días del periodo
    if dias_act is None:
        dias_act = (res["fecha_er"] - date(res["fecha_er"].year, 1, 1)).days + 1 if res["fecha_er"] else 365
    dias_ant = dias_ant or dias_act
    out["dias"] = [dias_act, dias_ant]
    meta["dias"] = ("Días desde el 1 de enero hasta la fecha de corte. Confirma que la columna anterior "
                    "sea el MISMO periodo del año previo.") if res["fecha_er"] else "Se asumió año completo (365)"

    # ---- validaciones (en millones)
    val = []

    def chk(nombre, a, b, tol=0.5):
        if a is None or b is None:
            val.append((nombre, None, None, "No evaluable"))
        else:
            ok = abs(a - b) <= max(tol, 0.0005 * max(abs(a), abs(b)))
            val.append((nombre, round(a, 2), round(b, 2), "OK" if ok else "REVISAR"))

    for etiqueta, idx in (("actual", ia), ("anterior", ip)):
        ta, tp, pat = get("total_activos", idx), get("total_pasivos", idx), get("patrimonio", idx)
        chk(f"Activos = Pasivos + Patrimonio ({etiqueta})", ta, (tp + pat) if tp is not None and pat is not None else None)
        if "total_ac" in d and "total_anc" in d:
            chk(f"Total activos = AC + ANC ({etiqueta})", ta, get("total_ac", idx) + get("total_anc", idx))
        ub, ing, cos = get("utilidad_bruta", idx), get("ingresos", idx), get("costo", idx)
        if ub is not None and ing is not None and cos is not None:
            chk(f"Utilidad bruta = Ingresos − Costo ({etiqueta})", ub, ing - abs(cos))
        un = get("utilidad_neta", idx)
        if un is not None and None not in (ing, cos):
            ub_ = ing - abs(cos)
            calc = ub_ - (out["gastos_op"][0 if idx == ia else 1] or 0) - (out["da"][0 if idx == ia else 1] or 0) \
                   - (out["gastos_fin"][0 if idx == ia else 1] or 0) - (out["impuestos"][0 if idx == ia else 1] or 0)
            chk(f"Utilidad neta reportada = calculada por la plantilla ({etiqueta})", un, calc)
    res["validaciones"] = val
    res["salida"] = out
    res["meta"] = meta
    res["unidad_aplicada"] = res["unidad"] if factor_override is None else "según --unidad"
    return out, meta


# --------------------------------------------------------------------------------------
# 7. NOTAS: segmentación y categorías
# --------------------------------------------------------------------------------------
def segmentar_notas(lineas):
    notas, actual = [], None
    pat_nota = re.compile(r"^\s*nota\s+(\d{1,2})\s*[\.\-:]?\s*(.*)$", re.I)
    pat_num = re.compile(r"^\s*(\d{1,2})[\.\)]\s+([A-ZÁÉÍÓÚÑ][^\d]{3,80})$")
    for ln in lineas:
        m = pat_nota.match(ln) or pat_num.match(ln)
        if m:
            actual = dict(nota=int(m.group(1)), titulo=m.group(2).strip(), texto=[])
            notas.append(actual)
        elif actual is not None:
            actual["texto"].append(ln)
        else:
            if not notas:
                actual = dict(nota=0, titulo="Introducción", texto=[ln])
                notas.append(actual)
    for n in notas:
        n["texto"] = " ".join(n["texto"])
        t = norm(n["titulo"] + " " + n["texto"])
        n["categorias"] = [c for c, kws in CATEGORIAS_NOTAS.items() if any(k in t for k in kws)]
    return notas


# --------------------------------------------------------------------------------------
# 8. ESCRIBIR EN LA PLANTILLA
# --------------------------------------------------------------------------------------
NARANJA = PatternFill("solid", start_color="FED7AA", end_color="FED7AA")
ROJO = PatternFill("solid", start_color="FECACA", end_color="FECACA")
AZUL_HDR = PatternFill("solid", start_color="0E7490", end_color="0E7490")


def escribir_plantilla(res, notas, ruta_plantilla, ruta_salida, adic_filas=None):
    wb = load_workbook(ruta_plantilla)
    ws = wb["Datos"]
    out, meta = res["salida"], res["meta"]
    ws["A2"] = (f"Cliente: extraído de PDF · Unidad original: {res['unidad'] or 'no detectada'} · "
                f"Cifras en millones de COP · Revisar hoja 'Revisión'")
    for clave, fila in FILAS_DATOS.items():
        for col, idx in ((2, 0), (3, 1)):
            celda = ws.cell(row=fila, column=col)
            v = out.get(clave, [None, None])[idx]
            if v is None:
                celda.value = None
                celda.fill = ROJO
            else:
                celda.value = round(v, 2) if clave != "dias" else int(v)
                if clave in ("otros_ac", "otros_pc", "gastos_op", "servicio_deuda", "dias") or \
                        (clave in meta and "confianza" in meta[clave]) or (clave in meta and "OCR" in meta[clave]):
                    celda.fill = NARANJA
                else:
                    celda.fill = PatternFill("solid", start_color="FFF9C4", end_color="FFF9C4")
        if clave in meta:
            ws.cell(row=fila, column=2).comment = Comment(meta[clave], "Valorix360")
    ws.cell(row=2, column=1).font = Font(name="Arial", size=10, italic=True, color="475569")

    # Datos que el PDF no trae: quitar los ejemplos de la plantilla y dejarlos para digitar
    for fila, _ in FILAS_MANUALES_CERO.items():
        for col in (2, 3):
            ws.cell(row=fila, column=col).value = 0
    for fila, _ in FILAS_MANUALES_VACIAS.items():
        celda = ws.cell(row=fila, column=2)
        celda.value = None
        celda.fill = ROJO
    for fila, valor in (adic_filas or {}).items():  # supuestos digitados en la app
        if valor is not None:
            celda = ws.cell(row=fila, column=2)
            celda.value = valor
            celda.fill = PatternFill("solid", start_color="FFF9C4", end_color="FFF9C4")
    if "Histórico" in wb.sheetnames:  # quitar los datos de ejemplo del histórico
        hs = wb["Histórico"]
        for fila in range(6, 15):
            for col in (2, 3, 4):
                hs.cell(row=fila, column=col).value = None

    # ---- Hoja Revisión
    for nombre in ("Revisión", "Notas"):
        if nombre in wb.sheetnames:
            del wb[nombre]
    rv = wb.create_sheet("Revisión")
    rv.sheet_view.showGridLines = False
    rv["A1"] = "Revisión de la extracción"
    rv["A1"].font = Font(name="Arial", size=14, bold=True)
    fila = 3

    def titulo(txt):
        nonlocal fila
        c = rv.cell(row=fila, column=1, value=txt)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        for col in range(1, 8):
            rv.cell(row=fila, column=col).fill = PatternFill("solid", start_color="0B1220", end_color="0B1220")
        fila += 1

    def encabezado(vals):
        nonlocal fila
        for i, v in enumerate(vals, 1):
            c = rv.cell(row=fila, column=i, value=v)
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = AZUL_HDR
        fila += 1

    titulo("AVISOS")
    avisos = list(res["avisos"])
    avisos.append(f"Periodos detectados (columnas): {res['years']} · actual = "
                  f"{res['years'][res['i_act']] if res['years'][0] else 'col. 1'}")
    if res["fecha_er"]:
        avisos.append(f"Fecha de corte del estado de resultados: {res['fecha_er']} -> días del periodo "
                      f"{out['dias'][0]}. Si el corte es intermedio, compara contra el MISMO periodo del año anterior.")
    for a in avisos:
        rv.cell(row=fila, column=1, value=a).font = Font(name="Arial", size=10)
        fila += 1
    fila += 1

    titulo("VALIDACIONES CONTABLES (millones COP)")
    encabezado(["Validación", "Valor A", "Valor B", "Estado"])
    for nombre, a, b, est in res["validaciones"]:
        rv.cell(row=fila, column=1, value=nombre)
        rv.cell(row=fila, column=2, value=a)
        rv.cell(row=fila, column=3, value=b)
        c = rv.cell(row=fila, column=4, value=est)
        c.font = Font(name="Arial", bold=True, color="166534" if est == "OK" else "991B1B")
        fila += 1
    fila += 1

    titulo("DATOS CARGADOS EN LA PLANTILLA")
    encabezado(["Dato", "Actual", "Anterior", "Origen / observación"])
    for clave, fila_d in FILAS_DATOS.items():
        rv.cell(row=fila, column=1, value=ws.cell(row=fila_d, column=1).value)
        rv.cell(row=fila, column=2, value=out.get(clave, [None, None])[0])
        rv.cell(row=fila, column=3, value=out.get(clave, [None, None])[1])
        rv.cell(row=fila, column=4, value=meta.get(clave, "Directo del PDF" if clave in res["datos"] else "FALTA: digitar manualmente"))
        fila += 1
    fila += 1

    titulo("DATOS PARA DIGITAR A MANO (hoja Datos, celdas en rojo)")
    encabezado(["Dato", "Fila en Datos", "", "Para qué se usa"])
    usos = {47: "LTM", 48: "LTM", 49: "LTM", 50: "LTM", 51: "LTM", 52: "LTM",
            54: "Covenants (nota de obligaciones)", 55: "Covenants", 56: "Covenants",
            57: "Perfil de vencimientos", 58: "Perfil de vencimientos", 59: "Perfil de vencimientos",
            60: "Perfil de vencimientos", 63: "WACC y EVA (hoja Valor y riesgo)"}
    for fila_m, nombre_m in FILAS_MANUALES_VACIAS.items():
        rv.cell(row=fila, column=1, value=nombre_m)
        rv.cell(row=fila, column=2, value=fila_m)
        rv.cell(row=fila, column=4, value=usos.get(fila_m, ""))
        fila += 1
    rv.cell(row=fila, column=1, value="Gastos e ingresos no recurrentes quedan en 0: ajústalos si las notas revelan partidas únicas.")
    fila += 2

    titulo("LÍNEAS DEL PDF SIN MAPEAR (revisar si alguna es relevante)")
    encabezado(["Etiqueta", "Valor 1", "Valor 2", "Sección", "Grupo", "Página"])
    for r in res["sin_mapear"]:
        rv.cell(row=fila, column=1, value=r["etiqueta"])
        rv.cell(row=fila, column=2, value=r["valores"][0])
        rv.cell(row=fila, column=3, value=r["valores"][1])
        rv.cell(row=fila, column=4, value=r["seccion"])
        rv.cell(row=fila, column=5, value=r["grupo"])
        rv.cell(row=fila, column=6, value=r["pagina"])
        fila += 1
    fila += 1

    titulo("LÍNEAS MAPEADAS (trazabilidad)")
    encabezado(["Etiqueta en el PDF", "Valor 1", "Valor 2", "Clave", "Confianza", "Método", "Página"])
    for r in res["log"]:
        for col, v in enumerate([r["etiqueta"], r["valores"][0], r["valores"][1], r["clave"],
                                 r["confianza"], r["metodo"], r["pagina"]], 1):
            rv.cell(row=fila, column=col, value=v)
        fila += 1
    for col, w in zip("ABCDEFG", (62, 16, 16, 40, 12, 22, 9)):
        rv.column_dimensions[col].width = w
    for row in rv.iter_rows():
        for c in row:
            if c.font is None or c.font.name != "Arial":
                c.font = Font(name="Arial", size=10, bold=c.font.bold if c.font else False,
                              color=c.font.color if c.font else None)

    # ---- Hoja Notas
    nt = wb.create_sheet("Notas")
    nt.sheet_view.showGridLines = False
    nt["A1"] = "Notas a los estados financieros (segmentadas y clasificadas)"
    nt["A1"].font = Font(name="Arial", size=14, bold=True)
    for i, h in enumerate(["Nota", "Título", "Temas detectados", "Texto"], 1):
        c = nt.cell(row=3, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = AZUL_HDR
    for r, n in enumerate(notas, 4):
        nt.cell(row=r, column=1, value=n["nota"])
        nt.cell(row=r, column=2, value=n["titulo"])
        nt.cell(row=r, column=3, value=", ".join(n["categorias"]))
        c = nt.cell(row=r, column=4, value=n["texto"][:1800])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        nt.row_dimensions[r].height = 75
    for col, w in zip("ABCD", (7, 34, 36, 110)):
        nt.column_dimensions[col].width = w
    wb.save(ruta_salida)


# --------------------------------------------------------------------------------------
# 8b. INDICADORES (misma lógica que la hoja "Indicadores" de la plantilla; para la app web)
# --------------------------------------------------------------------------------------
def calcular_indicadores(out):
    """out: {clave: [actual, anterior]} en millones. Devuelve lista de dicts con semáforo."""
    def v(k, i):
        x = out.get(k, [None, None])[i]
        return float(x) if x is not None else 0.0

    def base(i):
        ac = v("efectivo", i) + v("cxc", i) + v("inventarios", i) + v("otros_ac", i)
        pc = v("oblig_cp", i) + v("proveedores", i) + v("otros_pc", i)
        act, pas = ac + v("anc", i), pc + v("pnc", i)
        ub = v("ingresos", i) - v("costo", i)
        ebitda = ub - v("gastos_op", i)
        uo = ebitda - v("da", i)
        un = uo - v("gastos_fin", i) - v("impuestos", i)
        return dict(ac=ac, pc=pc, act=act, pas=pas, ub=ub, ebitda=ebitda, uo=uo, un=un,
                    dias=v("dias", i) or 365.0, fcl=v("cfo", i) - v("capex", i), i=i)

    def d(a, b):
        return a / b if b else 0.0

    ind = [  # grupo, nombre, función, referencia, mejor si
        ("Liquidez", "Razón corriente", lambda b: d(b["ac"], b["pc"]), 1.5, "Mayor"),
        ("Liquidez", "Prueba ácida", lambda b: d(b["ac"] - v("inventarios", b["i"]), b["pc"]), 1.0, "Mayor"),
        ("Liquidez", "Capital de trabajo neto ($M)", lambda b: b["ac"] - b["pc"], 0.0, "Mayor"),
        ("Liquidez", "Razón de efectivo", lambda b: d(v("efectivo", b["i"]), b["pc"]), 0.2, "Mayor"),
        ("Flujo de caja", "Calidad de utilidades", lambda b: d(v("cfo", b["i"]), b["un"]), 1.0, "Mayor"),
        ("Flujo de caja", "Flujo de caja libre ($M)", lambda b: b["fcl"], 0.0, "Mayor"),
        ("Flujo de caja", "Cobertura del servicio de deuda", lambda b: d(v("cfo", b["i"]), v("servicio_deuda", b["i"])), 1.25, "Mayor"),
        ("Flujo de caja", "Días de caja",
         lambda b: d(v("efectivo", b["i"]), (v("costo", b["i"]) + v("gastos_op", b["i"])) / b["dias"]), 30.0, "Mayor"),
        ("Utilidades", "Margen bruto", lambda b: d(b["ub"], v("ingresos", b["i"])), 0.30, "Mayor"),
        ("Utilidades", "Margen EBITDA", lambda b: d(b["ebitda"], v("ingresos", b["i"])), 0.15, "Mayor"),
        ("Utilidades", "Margen neto", lambda b: d(b["un"], v("ingresos", b["i"])), 0.08, "Mayor"),
        ("Utilidades", "ROE (anualizado)", lambda b: d(b["un"], v("patrimonio", b["i"])) * 365 / b["dias"], 0.12, "Mayor"),
        ("Utilidades", "ROA (anualizado)", lambda b: d(b["un"], b["act"]) * 365 / b["dias"], 0.06, "Mayor"),
        ("Inventarios y ciclo", "Rotación de inventarios (anualizada)",
         lambda b: d(v("costo", b["i"]), v("inventarios", b["i"])) * 365 / b["dias"], 6.0, "Mayor"),
        ("Inventarios y ciclo", "Días de inventario",
         lambda b: d(v("inventarios", b["i"]), v("costo", b["i"])) * b["dias"], 60.0, "Menor"),
        ("Inventarios y ciclo", "Días de cartera",
         lambda b: d(v("cxc", b["i"]), v("ingresos", b["i"])) * b["dias"], 45.0, "Menor"),
        ("Inventarios y ciclo", "Días de proveedores",
         lambda b: d(v("proveedores", b["i"]), v("costo", b["i"])) * b["dias"], 45.0, "Mayor"),
        ("Inventarios y ciclo", "Ciclo de conversión de efectivo",
         lambda b: d(v("inventarios", b["i"]), v("costo", b["i"])) * b["dias"]
         + d(v("cxc", b["i"]), v("ingresos", b["i"])) * b["dias"]
         - d(v("proveedores", b["i"]), v("costo", b["i"])) * b["dias"], 60.0, "Menor"),
        ("Endeudamiento", "Endeudamiento total", lambda b: d(b["pas"], b["act"]), 0.60, "Menor"),
        ("Endeudamiento", "Cobertura de intereses", lambda b: d(b["uo"], v("gastos_fin", b["i"])), 3.0, "Mayor"),
    ]
    b0, b1 = base(0), base(1)
    filas = []
    for grupo, nombre, fn, ref, mejor in ind:
        a, p = fn(b0), fn(b1)
        ok = a >= ref if mejor == "Mayor" else a <= ref
        filas.append({"Grupo": grupo, "Indicador": nombre, "Actual": round(a, 3), "Anterior": round(p, 3),
                      "Referencia": ref, "Mejor si": mejor, "Estado": "Bien" if ok else "Revisar"})
    return filas


# --------------------------------------------------------------------------------------
# 9. CLI
# --------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Extractor de estados financieros PDF -> plantilla Valorix360")
    ap.add_argument("pdf")
    ap.add_argument("--plantilla", default="valorix360_plantilla_diagnostico.xlsx")
    ap.add_argument("--salida", default="diagnostico_cliente.xlsx")
    ap.add_argument("--unidad", choices=["pesos", "miles", "millones"], help="Forzar la unidad del PDF")
    ap.add_argument("--dias-actual", type=int)
    ap.add_argument("--dias-anterior", type=int)
    ap.add_argument("--da-no-incluida", action="store_true",
                    help="La depreciación NO está dentro de los gastos operativos (no restarla)")
    ap.add_argument("--json-notas", help="Guardar las notas segmentadas en este archivo JSON")
    a = ap.parse_args()

    res = extraer(a.pdf)
    f = {"pesos": 1e-6, "miles": 1e-3, "millones": 1.0}.get(a.unidad)
    consolidar(res, f, a.dias_actual, a.dias_anterior, da_en_gastos=not a.da_no_incluida)
    notas = segmentar_notas(res["notas_lineas"])
    escribir_plantilla(res, notas, a.plantilla, a.salida)
    if a.json_notas:
        with open(a.json_notas, "w", encoding="utf-8") as fh:
            json.dump(notas, fh, ensure_ascii=False, indent=2)

    print(f"Unidad detectada: {res['unidad']} · periodos: {res['years']} · notas: {len(notas)}")
    malas = [v for v in res["validaciones"] if v[3] == "REVISAR"]
    print(f"Validaciones: {len(res['validaciones']) - len(malas)} OK, {len(malas)} por revisar")
    for av in res["avisos"]:
        print("AVISO:", av)
    faltan = [k for k in FILAS_DATOS if res["salida"].get(k, [None])[0] is None]
    if faltan:
        print("Datos que faltan (digitar a mano):", ", ".join(faltan))
    print(f"Listo -> {a.salida}. Abre la hoja 'Revisión' antes de usar los resultados.")


if __name__ == "__main__":
    sys.exit(main())
