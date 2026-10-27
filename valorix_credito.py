"""Adaptador y perfil crediticio interno.

El perfil no se presenta en la interfaz del cliente. Sus resultados son indicativos,
explicables y requieren validación humana antes de solicitar crédito.
"""
import io
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from valorix_motor import ADIC_DEFECTO, periodo

PLAZOS_CREDITO = (12, 24, 36)
DSCR_RECOMENDADO = 1.25
DSCR_MAXIMO = 1.10
MAX_MB = 25

_COLUMNAS = (3, 5, 7, 9)
_PYG = {
    6: "ingresos", 8: "costo", 10: "gastos_op", 16: "impuestos",
    17: "gastos_fin",
}
_BALANCE = {
    23: "efectivo", 24: "cxc", 25: "inventarios", 26: "otros_ac", 29: "anc",
    35: "oblig_cp", 36: "proveedores", 37: "otros_pc", 40: "deuda_lp",
    41: "pnc", 49: "patrimonio", 45: "capital",
}


class ErrorVaciado(ValueError):
    """Libro interno inválido o incompleto."""


def _valor(celda):
    if isinstance(celda, bool) or isinstance(celda, str):
        return None
    try:
        return None if celda is None else float(celda)
    except (TypeError, ValueError):
        return None


def _numero(valor, defecto=0.0):
    return float(valor) if isinstance(valor, (int, float)) and not isinstance(valor, bool) else defecto


def es_vaciado_credito(contenido):
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError):
        return False
    try:
        return "Vaciado de CIFRAS" in wb.sheetnames
    finally:
        wb.close()


def cargar_vaciado_credito(contenido, nombre="vaciado.xlsx"):
    """Convierte el libro interno a datos del diagnóstico y periodos crediticios."""
    if len(contenido) > MAX_MB * 1024 * 1024:
        raise ErrorVaciado(f"El archivo supera el límite de {MAX_MB} MB.")
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as e:
        raise ErrorVaciado("No se pudo leer el vaciado interno.") from e
    if "Vaciado de CIFRAS" not in wb.sheetnames:
        wb.close()
        raise ErrorVaciado("El libro no contiene la hoja 'Vaciado de CIFRAS'.")
    try:
        ws = wb["Vaciado de CIFRAS"]
        periodos = []
        for col in _COLUMNAS:
            periodo_nombre = ws.cell(4, col).value
            dias = _valor(ws.cell(5, col).value)
            salida: dict[str, list[float | None]] = {k: [None, None] for k in (
                "efectivo", "cxc", "inventarios", "otros_ac", "anc", "oblig_cp",
                "proveedores", "otros_pc", "pnc", "patrimonio", "ingresos", "costo",
                "gastos_op", "da", "gastos_fin", "impuestos", "cfo", "capex",
                "servicio_deuda", "dias", "deuda_lp", "capital",
            )}
            salida["dias"] = [dias, dias]
            for row, clave in {**_PYG, **_BALANCE}.items():
                salida[clave][0] = _valor(ws.cell(row, col).value)
            salida["gastos_op"][0] = sum(
                v for v in (_valor(ws.cell(10, col).value), _valor(ws.cell(11, col).value))
                if v is not None
            ) or None
            utilidad = _valor(ws.cell(18, col).value)
            if utilidad is None and salida["ingresos"][0] is not None:
                utilidad = salida["ingresos"][0] - sum(
                    _numero(salida[k][0]) for k in ("costo", "gastos_op", "gastos_fin", "impuestos")
                )
            salida["utilidad_neta"] = [utilidad, None]
            if salida["patrimonio"][0] is None:
                componentes = [_valor(ws.cell(row, col).value) for row in (45, 46, 47)]
                if utilidad is not None and all(v is not None for v in componentes):
                    salida["patrimonio"][0] = sum(_numero(v) for v in componentes) + utilidad
            periodos.append(dict(nombre=periodo_nombre, dias=dias, salida=salida,
                                 disponible=all(salida[k][0] is not None for k in ("ingresos", "efectivo", "patrimonio"))))
    finally:
        wb.close()
    completos = [p for p in periodos if p["disponible"]]
    if not completos:
        raise ErrorVaciado("El vaciado no tiene un periodo con cifras mínimas.")
    actual = completos[-1]["salida"]
    anterior = completos[-2]["salida"] if len(completos) > 1 else None
    for clave in actual:
        actual[clave][1] = None if anterior is None else anterior[clave][0]
    return dict(datos={}, log=[], sin_mapear=[], avisos=[
        f"Vaciado interno cargado desde {nombre}; el perfil crediticio no se muestra al cliente."
    ], factor=1.0, unidad="millones de COP", i_act=0, years=[],
                fecha_er=None, notas_lineas=[], salida=actual, meta={},
                validaciones=[], periodos=periodos)


def _flujo_disponible(out):
    p = periodo(out, 0)
    cfo = out.get("cfo", [None])[0]
    capex = out.get("capex", [None])[0]
    if cfo is not None and capex is not None:
        return max(0.0, _numero(cfo) - _numero(capex)), "flujo de caja libre"
    proxy = max(0.0, _numero(p.get("ebitda")) - _numero(p.get("da"))
                - _numero(p.get("gf")) - _numero(p.get("imp")))
    return proxy, "proxy de flujo (faltan CFO y Capex)"


def _valor_presente(cuota_anual, tasa_anual, meses):
    r = max(0.0, tasa_anual) / 12
    if r == 0:
        return cuota_anual * meses / 12
    return cuota_anual / 12 * (1 - (1 + r) ** -meses) / r


def recomendar_credito(out, adic=None, tasa_anual=None):
    """Calcula montos indicativos con y sin deuda vigente para cada plazo."""
    adic = dict(ADIC_DEFECTO, **(adic or {}))
    p = periodo(out, 0)
    flujo, fuente_flujo = _flujo_disponible(out)
    deuda_actual = max(0.0, _numero(p.get("sd")))
    if deuda_actual == 0:
        deuda_actual = max(0.0, _numero(p.get("gf")) * 1.2)
    deuda_base = max(0.0, _numero(p.get("dcp")) + _numero(p.get("dlp")))
    tasa = tasa_anual
    if tasa is None:
        tasa = _numero(p.get("gf")) / deuda_base if deuda_base else 0.18
    tasa = min(max(float(tasa), 0.01), 0.60)
    escenarios = []
    for etiqueta, servicio_existente in (("con_deuda_actual", deuda_actual), ("sin_deuda_actual", 0.0)):
        cuotas = []
        for dscr, tipo in ((DSCR_RECOMENDADO, "recomendado"), (DSCR_MAXIMO, "maximo_indicativo")):
            servicio_nuevo = max(0.0, flujo / dscr - servicio_existente)
            for meses in PLAZOS_CREDITO:
                monto = _valor_presente(servicio_nuevo, tasa + (adic.get("est_tasa") or 0.0), meses)
                cuotas.append(dict(tipo=tipo, meses=meses, monto_m=round(monto, 2),
                                   cuota_anual_m=round(servicio_nuevo, 2), dscr=dscr))
        escenarios.append(dict(nombre=etiqueta, servicio_actual_m=round(servicio_existente, 2), cuotas=cuotas))
    return dict(version=1, visible_cliente=False, estado="indicativo",
                fuente_flujo=fuente_flujo, flujo_disponible_m=round(flujo, 2),
                tasa_anual=round(tasa, 6), tasa_con_estres=round(tasa + (adic.get("est_tasa") or 0.0), 6),
                dscr_recomendado=DSCR_RECOMENDADO, dscr_maximo=DSCR_MAXIMO,
                plazos=list(PLAZOS_CREDITO), deuda_actual_m=round(deuda_actual, 2),
                escenarios=escenarios,
                advertencias=["No constituye aprobación bancaria ni asesoría financiera.",
                              "Debe validarse con estados, deuda y flujo de caja completos."])
