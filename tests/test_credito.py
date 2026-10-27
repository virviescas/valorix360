import io

from openpyxl import Workbook

from valorix_credito import (DSCR_MAXIMO, DSCR_RECOMENDADO, cargar_vaciado_credito,
                             recomendar_credito)


def _libro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Vaciado de CIFRAS"
    ws["C4"], ws["E4"], ws["G4"], ws["I4"] = 2024, 2025, "Ago 2025", "Ago 2026"
    ws["C5"], ws["E5"], ws["G5"], ws["I5"] = 360, 360, 240, 240
    for col in (3, 5, 9):
        ws.cell(6, col).value = 1000
        ws.cell(23, col).value = 100
        ws.cell(24, col).value = 100
        ws.cell(25, col).value = 100
        ws.cell(35, col).value = 40
        ws.cell(36, col).value = 60
        ws.cell(49, col).value = 300
    ws["C8"], ws["E8"], ws["I8"] = 500, 500, 500
    ws["C10"], ws["E10"], ws["I10"] = 100, 100, 100
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_vaciado_preserva_periodos_incompletos_y_no_usa_formulas():
    res = cargar_vaciado_credito(_libro())

    assert len(res["periodos"]) == 4
    assert res["periodos"][2]["disponible"] is False
    assert res["salida"]["ingresos"] == [1000.0, 1000.0]


def test_recomendacion_escalonada_por_dscr_y_plazo():
    out = cargar_vaciado_credito(_libro())["salida"]
    out["cfo"] = [300, 300]
    out["servicio_deuda"] = [60, 60]
    perfil = recomendar_credito(out, {"est_tasa": 0.03})

    assert perfil["visible_cliente"] is False
    assert perfil["dscr_recomendado"] == DSCR_RECOMENDADO
    assert perfil["dscr_maximo"] == DSCR_MAXIMO
    assert perfil["plazos"] == [12, 24, 36]
    assert len(perfil["escenarios"]) == 2
    recomendado = perfil["escenarios"][0]["cuotas"]
    assert recomendado[0]["monto_m"] < recomendado[3]["monto_m"]
